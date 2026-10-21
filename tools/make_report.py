# -*- coding: utf-8 -*-
"""Render report_content.CONTENT to Markdown and DOCX."""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from report_content import CONTENT, ROOT

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY = "宋体"
HEAD = "黑体"

# ---------------------------------------------------------------- markdown
md = []
for block in CONTENT:
    kind = block[0]
    if kind == "title":
        md.append("# " + block[1] + "\n")
    elif kind == "sub":
        md.append("> " + block[1] + "\n")
    elif kind == "h1":
        md.append("## " + block[1] + "\n")
    elif kind == "h2":
        md.append("### " + block[1] + "\n")
    elif kind == "p":
        md.append(block[1] + "\n")
    elif kind == "code":
        md.append("```c")
        md.extend(block[1])
        md.append("```\n")
    elif kind == "table":
        headers, rows = block[1], block[2]
        md.append("| " + " | ".join(headers) + " |")
        md.append("|" + "|".join([" --- "] * len(headers)) + "|")
        for r in rows:
            md.append("| " + " | ".join(r) + " |")
        md.append("")
    elif kind == "img":
        rel = os.path.relpath(block[1], ROOT).replace("\\", "/")
        md.append("![%s](%s)\n" % (block[2], rel))
        md.append("*%s*\n" % block[2])

io.open(os.path.join(ROOT, "实验报告.md"), "w", encoding="utf-8", newline="\n").write("\n".join(md))

# ---------------------------------------------------------------- docx
doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.3)
sec.bottom_margin = Cm(2.3)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)


def set_run(run, size=10.5, bold=False, font=BODY, ascii_font="Times New Roman", color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if color:
        run.font.color.rgb = color


def para(text, size=10.5, bold=False, indent=True, before=0, after=4, align=None,
         font=BODY, ascii_font="Times New Roman", spacing=1.35, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if align is not None:
        p.alignment = align
    set_run(p.add_run(text), size=size, bold=bold, font=font, ascii_font=ascii_font, color=color)
    return p


def heading(text, size, before=12, after=6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = True
    set_run(p.add_run(text), size=size, bold=True, font=HEAD)
    return p


def code_block(lines):
    for i, ln in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0 if i < len(lines) - 1 else 6)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.7)
        set_run(p.add_run(ln if ln else " "), size=9, font="Consolas", ascii_font="Consolas")
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "F2F4F7")
        p._p.get_or_add_pPr().append(shd)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(headers, rows, size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(h), size=size, bold=True)
        shade(cell, "DCE6F1")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            set_run(p.add_run(v), size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


for block in CONTENT:
    kind = block[0]
    if kind == "title":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        set_run(p.add_run(block[1]), size=19, bold=True, font=HEAD)
    elif kind == "sub":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        set_run(p.add_run(block[1]), size=9.5, color=RGBColor(0x55, 0x55, 0x55))
    elif kind == "h1":
        heading(block[1], 14, before=14, after=6)
    elif kind == "h2":
        heading(block[1], 11.5, before=8, after=4)
    elif kind == "p":
        para(block[1])
    elif kind == "code":
        code_block(block[1])
    elif kind == "table":
        table(block[1], block[2])
    elif kind == "img":
        doc.add_picture(block[1], width=Cm(7.6))
        pic_p = doc.paragraphs[-1]
        pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_p.paragraph_format.space_before = Pt(6)
        pic_p.paragraph_format.space_after = Pt(2)
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        set_run(cap.add_run(block[2]), size=9, color=RGBColor(0x40, 0x40, 0x40))

doc.save(os.path.join(ROOT, "实验报告.docx"))
print("md + docx written")
